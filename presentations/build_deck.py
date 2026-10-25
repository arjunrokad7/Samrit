"""Build Samrit's editable final presentation. Dependencies stay in workspace tools."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / '.tools' / 'presentation-python'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from PIL import Image
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
A = OUT / 'assets'
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
P.core_properties.title = 'Samrit — A place at the table'
P.core_properties.subject = 'Strict simplicity. Accountable food redistribution.'
P.core_properties.author = 'Samrit team'
P.core_properties.keywords = 'Samrit, Steel and Steam, food redistribution, UX, hackathon'
C = dict(ink='244557', teal='376C6F', fern='537A5A', spruce='2F4F31', iron='484D48',
         paper='F5F7F3', white='FFFFFF', mist='E4ECE7', pale='D9E3DB', lime='B5D588', dark='142A33', muted='B7CDD0')
HEAD, BODY = 'Arial Rounded MT Bold', 'Arial'
INDEX = []
SOURCES = [
 ('S1', 'USDA ERS · Food security, 2024', '47.9 million people lived in food-insecure U.S. households in 2024.', 'https://www.ers.usda.gov/topics/food-nutrition-assistance/food-security-in-the-us/key-statistics-graphics'),
 ('S2', 'EPA · U.S. food loss and waste', 'An estimated 30–40% of the U.S. food supply is never eaten.', 'https://www.epa.gov/sustainable-management-food/united-states-2030-food-loss-and-waste-reduction-goal'),
 ('S3', 'Food Bank For NYC · Distribution', 'Nearly 800 pantries, soup kitchens and mobile pantries; a mixed network, not a soup-kitchen-only count.', 'https://www.foodbanknyc.org/what-we-do/food-distribution/'),
 ('S4', 'Pew Research Center · Mobile, 2026', '80% smartphone ownership among U.S. adults in households earning below $30,000; survey Jan. 30–June 18, 2026.', 'https://www.pewresearch.org/internet/fact-sheet/mobile/'),
 ('S5', 'FDA · Nine major food allergens', 'Major-allergen categories inform the interface; the platform is not FDA endorsed.', 'https://www.fda.gov/food/buy-store-serve-safe-food/food-allergies-what-you-need-know'),
 ('S6', 'USDA · Donating / Good Samaritan Act', 'The Bill Emerson Good Samaritan Food Donation Act provides conditional protections for qualifying good-faith donations.', 'https://www.usda.gov/about-food/food-safety/food-loss-and-waste/donating'),
 ('S7', 'Cambridge Dictionary · Good Samaritan', 'A person who helps people in need; the name invokes this everyday meaning.', 'https://dictionary.cambridge.org/us/dictionary/english/good-samaritan'),
 ('S8', 'A Practical Sanskrit Dictionary · amṛta', 'Amrit derives from Sanskrit amṛta, including the sense nectar. Essence is the team’s brand interpretation.', 'https://api.pageplace.de/preview/DT0400.9788120837737_A24265371/preview-9788120837737_A24265371.pdf'),
 ('S9', 'OHCHR / FAO · The Right to Food', 'The right to adequate food is recognized in international human-rights instruments.', 'https://www.ohchr.org/sites/default/files/Documents/Publications/FactSheet34en.pdf'),
]

def color(v): return RGBColor.from_string(C.get(v, v))
def shape(slide, x, y, w, h, fill, rounded=False, line=None, kind=None):
    s = slide.shapes.add_shape(kind or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE), Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color(fill)
    if line: s.line.color.rgb = color(line); s.line.width = Pt(1)
    else: s.line.fill.background()
    if rounded:
        try: s.adjustments[0] = .12
        except Exception: pass
    return s
def text(slide, value, x, y, w, h, size=20, ink='ink', bold=False, font=BODY, align=None):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t = s.text_frame; t.clear(); t.word_wrap = True
    t.margin_left = t.margin_right = t.margin_top = t.margin_bottom = 0
    t.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(value.split('\n')):
        p = t.paragraphs[0] if i == 0 else t.add_paragraph()
        p.text = line; p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.08
        if align: p.alignment = align
        for r in p.runs:
            r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color(ink)
    return s
def rule(slide, x1, y1, x2, y2, ink='pale', width=1):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color(ink); s.line.width = Pt(width)
    return s
def logo(slide, x, y, w, light=False, leaf=False):
    path = ROOT / 'frontend/public/brand' / ('logoleaf.png' if leaf else ('samrit-light.png' if light else 'samrit-dark.png'))
    im = Image.open(path); iw, ih = im.size; l, t, r, b = im.getbbox(); h = w * (b-t)/(r-l)
    s = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    s.crop_left, s.crop_top, s.crop_right, s.crop_bottom = l/iw, t/ih, (iw-r)/iw, (ih-b)/ih
    return s
def base(kicker, title, dark=False, title_size=34):
    slide = P.slides.add_slide(P.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color('dark' if dark else 'paper')
    n = len(P.slides)
    text(slide, kicker.upper(), .65, .42, 10.6, .22, 11, 'lime' if dark else 'teal', True)
    if title: text(slide, title, .65, .96, 12, 1.6 if '\n' in title else 1.08, title_size, 'white' if dark else 'ink', font=HEAD)
    rule(slide, .65, 7.03, 12.68, 7.03, 'teal' if dark else 'pale')
    text(slide, 'SAMRIT  /  A PLACE AT THE TABLE', .65, 7.15, 8.5, .18, 8.5, 'muted' if dark else 'iron')
    text(slide, f'{n:02}', 12.15, 7.12, .5, .23, 11, 'muted' if dark else 'iron', align=PP_ALIGN.RIGHT)
    return slide
def notes(slide, title, body, refs=()):
    source_lines = '\n'.join(f'{k} · {next(s[3] for s in SOURCES if s[0] == k)}' for k in refs)
    slide.notes_slide.notes_text_frame.text = body + ('\n\nSources\n' + source_lines if refs else '')
    INDEX.append({'slide': len(P.slides), 'title': title, 'notes': body, 'sources': list(refs)})
def foot(slide, value, dark=False): text(slide, value, .65, 6.64, 12, .29, 10, 'muted' if dark else 'iron')
def link(slide, label, url, x, y, w, size=11, ink='teal'):
    s = text(slide, label, x, y, w, .25, size, ink)
    s.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    return s
def image(slide, path, x, y, w, h, frame=True, crop=None):
    path = Path(path)
    if frame:
        shape(slide, x+.035, y+.075, w, h, 'ink', True)
        shape(slide, x, y, w, h, 'white', True, 'pale')
    im = Image.open(path)
    iw, ih = im.size
    l,t,r,b = crop or (0,0,iw,ih)
    ratio = min((w-.12)/(r-l), (h-.12)/(b-t))
    pw, ph = (r-l)*ratio, (b-t)*ratio
    pic = slide.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l/iw,t/ih,(iw-r)/iw,(ih-b)/ih
    pic._element.nvPicPr.cNvPr.set('descr', 'Samrit interface screenshot or test fixture: ' + path.stem.replace('-', ' '))
    return pic
def point(slide, num, title, body, x, y, w=4, dark=False):
    text(slide, num, x, y, .45, .4, 17, 'lime' if dark else 'fern', True)
    text(slide, title, x+.58, y-.01, w-.58, .42, 21, 'white' if dark else 'ink', True)
    text(slide, body, x+.58, y+.46, w-.6, .88, 17, 'muted' if dark else 'iron')

# 01 · Name and promise
s = base('Steel & Steam · Food redistribution', '', True)
logo(s, .65, 1.05, 4.9, True)
text(s, 'A place at the table.\nA community behind it.', .65, 3.76, 7.15, 1.5, 34, 'white', font=HEAD)
text(s, 'Strict simplicity. Accountable food recovery.', .68, 5.66, 7, .45, 20, 'muted')
rule(s, 8.12, 1.23, 8.12, 6.3, 'teal', 1.5)
text(s, 'Samaritan', 8.65, 1.55, 3.8, .48, 29, 'white', font=HEAD)
text(s, 'Helping someone in need.\nA nod to the Bill Emerson\nGood Samaritan Food Donation Act.', 8.66, 2.15, 3.7, 1.15, 16, 'muted')
text(s, '+', 8.65, 3.47, 1, .5, 31, 'lime')
text(s, 'Amrit', 8.65, 4.14, 3.7, .5, 29, 'white', font=HEAD)
text(s, 'From Sanskrit amṛta: nectar.\nOur interpretation: the essence\nof nourishment and care.', 8.66, 4.78, 3.6, 1.2, 18, 'muted')
notes(s, 'Samrit', 'Samrit brings together the everyday idea of a Good Samaritan—a person helping someone in need—and Amrit, from Sanskrit amṛta, nectar. Essence is our brand interpretation rather than a literal dictionary gloss. The Samaritan reference also evokes the Bill Emerson Good Samaritan Food Donation Act. We are not claiming that using Samrit automatically provides immunity: donation protections have conditions. Our promise is a welcoming route to food, supported by an accountable community.', ['S6','S7','S8'])

# 02 · Kitchen interpretation of the theme
s = base('The theme, reinterpreted', 'Steel & steam.\nThe kitchen is our starting point.', True, 35)
point(s,'01','Steel is the structure.','Trucks, equipment, storage,\nrecords and reliable handoffs.',.7,2.62,6,True)
point(s,'02','Steam is the purpose.','A warm meal, a welcoming space,\nand the dignity of being nourished.',.7,4.48,6,True)
# Editable cooking-pot motif; no stock or generated image.
shape(s,8.16,3.62,3.28,1.7,'teal',True)
shape(s,7.72,3.75,.58,.32,'muted',True); shape(s,11.34,3.75,.58,.32,'muted',True)
shape(s,8.04,3.4,3.52,.16,'pale',True); shape(s,9.48,3.12,.64,.14,'pale',True)
for bx in [8.68,9.66,10.64]:
    pts=[(bx,2.98),(bx-.15,2.72),(bx+.1,2.43),(bx-.1,2.1)]
    for a,b in zip(pts,pts[1:]): rule(s,*a,*b,'lime',3)
rule(s,7.72,5.55,11.92,5.55,'fern',3)
text(s,'LOGISTICS IN SERVICE OF A MEAL',7.7,5.91,4.5,.5,13,'muted',True,align=PP_ALIGN.CENTER)
notes(s,'Steel & steam','The logistics interpretation of Steel and Steam is natural: vehicles, equipment and a chain of handoffs. Our twist is to start inside the kitchen. Steel is the physical and operational structure that makes a meal possible. Steam is the warm, human result. This is not a transport dashboard looking for a use case; it is a food-access experience with logistics behind it.')

# 03 · Carefully scoped statistics
s = base('The need and the opportunity', 'Food exists. Need exists.\nThe connection is missing.', False, 33)
stats = [('47.9M','PEOPLE','Lived in food-insecure\nU.S. households in 2024.','S1'),('30–40%','OF U.S. FOOD SUPPLY','Estimated to go uneaten.','S2'),('~800','NYC NETWORK LOCATIONS','Pantries, soup kitchens\nand mobile pantries.','S3'),('80%','SMARTPHONE OWNERSHIP','Adults in households\nbelow $30K income.\nUnited States, 2026.','S4')]
for i,(number,label,body,ref) in enumerate(stats):
    x=.65+i*3.1
    if i: rule(s,x-.2,2.43,x-.2,5.9,'pale')
    text(s,number,x,2.6,2.85,.85,43,'teal',font=HEAD)
    text(s,label,x,3.72,2.83,.48,11,'ink',True)
    text(s,body,x,4.39,2.73,1.05,18,'iron')
    src=next(v for v in SOURCES if v[0]==ref)
    link(s,f'[{ref}] {src[1].split(" · ")[0]}',src[3],x,5.64,2.8,10)
foot(s,'Phone ownership is an access proxy, not a measurement of food-insecure people. NYC network count combines service types.')
notes(s,'Need and access','These figures describe different populations and should not be multiplied together. USDA reports 47.9 million people living in food-insecure U.S. households in 2024; this does not mean every person experienced hunger in the same way. EPA estimates 30–40% of the U.S. food supply goes uneaten. Food Bank For NYC describes a network of nearly 800 pantries, soup kitchens and mobile pantries; it is not a soup-kitchen-only census or a Samrit partner count. Pew’s 2026 survey found 80% smartphone ownership among adults in households earning under $30,000. That is a proxy for reach, not a food-insecurity-specific smartphone measure. Digital access still has gaps.', ['S1','S2','S3','S4'])

# 04 · Principles
s = base('Our design principles', 'Strict simplicity.\nAccountable simplicity.', True, 43)
for i,(title,body) in enumerate([('People seeking food first.','No account required to find meals.\nThe primary action is easy to see.'),('Accessible by intention.','Plain language, responsive screens,\nvisible choices and optional voice.'),('Simple does not mean opaque.','A clear next step for the person.\nA clear record for the operator.')]):
    x=.65+i*4.13
    rule(s,x,3.12,x+3.72,3.12,'fern',3)
    text(s,f'0{i+1}',x,3.47,1,.6,32,'lime',font=HEAD)
    text(s,title,x,4.35,3.7,.8,23,'white',True)
    text(s,body,x,5.38,3.7,1.0,17,'muted')
notes(s,'Strict simplicity','The person looking for food should not have to understand the logistics system. Our landing page prioritizes finding a meal. A person can browse without an account, choose visible allergen chips, write preferences in everyday language, or use the optional assistant. Staff workspaces are focused on their own jobs. Accountability lives in role permissions, verified menus, stage-ordered handoffs and evidence records. Simplicity is the front door; accountability supports it behind the scenes. These are design goals, not a claim of a completed formal accessibility audit.')

# 05 · What distinguishes Samrit
s = base('What makes Samrit different', 'Four capabilities. One coherent experience.')
features=[('01','Dignified discovery','Allergen and preference filtering,\nwith optional voice assistance.'),('02','Menus into structured data','Gemini OCR creates editable dish\ncandidates for human confirmation.'),('03','Traceable handoffs','QR events, photo hashes and\nlocation provenance in one record.'),('04','Private quality learning','Internal feedback helps the NGO\nguide donors and venues.')]
for i,(num,title,body) in enumerate(features):
    x=.65+(i%2)*6.18; y=2.25+(i//2)*2.1
    shape(s,x+.035,y+.07,5.82,1.76,'spruce',True)
    shape(s,x,y,5.82,1.76,'white',True,'pale')
    text(s,num,x+.22,y+.27,.55,.5,24,'fern',font=HEAD)
    text(s,title,x+.94,y+.25,4.6,.52,23,'ink',True)
    text(s,body,x+.94,y+.94,4.55,.66,17,'iron')
notes(s,'Unique capabilities','The platform’s distinctive contribution is the combination: accessible discovery, human-verified menu extraction, accountable handoffs and private organizational learning. The AI is deliberately connected to concrete actions and structured records. It is not used as a replacement for food-handler judgment or as a public score that can shame a venue. The next slides show how these capabilities are distributed across the four faces.')

# 06 · Four faces
s = base('The platform structure', 'Four faces. A shared responsibility.', True)
roles=[('01','People seeking\nmeals','Discover, choose,\ncheck in and review.'),('02','Distribution\ncenters','Profile capacity,\nreceive and verify.'),('03','Donors &\norganizers','Confirm menus,\noffer leftovers.'),('04','Central NGO','Approve, dispatch,\nreview and guide.')]
for i,(num,title,body) in enumerate(roles):
    x=.65+3.1*i
    shape(s,x,2.52,2.78,3.51,'teal' if i==0 else 'ink',True)
    text(s,num,x+.23,2.82,1,.65,34,'lime',font=HEAD)
    text(s,title,x+.23,3.79,2.38,.96,23,'white',True)
    text(s,body,x+.23,5.12,2.4,.7,17,'muted')
foot(s,'The NGO face includes its driver mobile view. Drivers are provisioned by the NGO, not recruited through public signup.',True)
notes(s,'Four-face structure','We retain the original four-face structure. Face one is the public consumer experience; face two is the distribution center; face three is the event organizer or donor; and face four is the central NGO. The driver mobile view belongs to the NGO logistics face rather than becoming a fifth public signup category. The NGO already manages its people and fleet. Samrit coordinates those resources and records their handoffs.')

# 07 · Consumer
s = base('Face 01 · People seeking meals', 'Find a place that works for you.')
point(s,'01','Choose what fits.','Nine major-allergen categories\nand natural-language preferences.',.65,2.35,4.1)
point(s,'02','Stay in control.','Visible selections, optional voice\nand no public account requirement.',.65,3.8,4.1)
point(s,'03','Be heard privately.','One venue QR, finish the visit,\nthen leave private feedback.',.65,5.24,4.1)
image(s,A/'consumer.png',5.1,2.05,7.55,4.75)
notes(s,'Consumer face','The person seeking a meal is the priority, not a side link inside a staff portal. The interface uses FDA-recognized major-allergen categories: milk, eggs, fish, Crustacean shellfish, tree nuts, peanuts, wheat, soybeans and sesame. Gemini can interpret a preference such as avoiding cold food, while unknown nutritional properties are not invented. People should still confirm ingredients and cross-contact with staff. The sourced NYC directory and fictional demo meal catalog are distinct; this is not a live promise of meal availability. One shared venue QR starts the visit, and a finish button opens private feedback without an exit scan.', ['S5'])

# 08 · Distribution center
s = base('Face 02 · Distribution centers', 'A welcoming place, ready to receive.')
image(s,A/'center.png',.65,2.13,7.62,4.61)
point(s,'01','Capacity first.','Hours, seating, staffing and\nhot / chilled / frozen capability.',8.7,2.45,3.95)
point(s,'02','NGO-approved.','The central NGO verifies\nparticipating facilities.',8.7,3.94,3.95)
point(s,'03','Staff confirm intake.','A shared venue QR and staff\narrival-photo confirmation.',8.7,5.39,3.95)
notes(s,'Distribution-center face','A center registers its facility, operating hours, equipment, seating, staffing and receiving capacity. Admin approval makes the center eligible for matching. Mixed batches require support for every included handling type. The driver scans the center’s shared venue code on arrival; center staff scan the tracking sticker and submit the receiving photograph. The same venue code supports consumer check-in, reducing signage and removing an unnecessary second entry code and exit code. A shared code does not merge delivery records with anonymous visit credentials.')

# 09 · Donor / OCR
s = base('Face 03 · Event organizers & donors', 'From a menu photo to a rescue batch.')
image(s,A/'sample-menu.png',.65,2.12,3.45,3.98)
image(s,A/'donor.png',4.48,2.12,8.17,3.98)
for i,(label,sub) in enumerate([('EXTRACT','Gemini reads the menu / invoice.'),('CONFIRM','The donor verifies every dish.'),('SELECT LEFTOVERS','Positive servings + handling types.')]):
    x=.75+4.16*i
    text(s,label,x,6.29,3.8,.23,12,'teal',True)
    text(s,sub,x,6.62,3.9,.23,11,'iron')
notes(s,'Donor and OCR','The donor can upload a menu or invoice for Gemini extraction. Dishes, ingredients, declared allergens, servings and handling requirements become editable candidates. The donor must confirm them; an invoice’s ingredients do not prove that a dish was served. After the event starts, the donor selects leftovers, enters positive whole-number servings, chooses all needed handling types and adds the Stage 1 photo. Clearing a servings field keeps its checkbox selected and blocks submission until corrected. Mixed handling is supported only when the assigned vehicle and destination can support every included requirement. The menu shown is a fictional test fixture.')

# 10 · NGO / driver
s = base('Face 04 · Central NGO + driver view', 'The right food. The right destination.')
point(s,'01','Existing resources.','Approve venues; provision drivers;\nmaintain vehicles and equipment.',.65,2.32,4.15)
point(s,'02','Compatible dispatch.','Match total capacity, handling,\nopening hours and food expiry.',.65,3.84,4.15)
point(s,'03','Clear driver steps.','Event QR → pickup photo → sticker\n→ venue QR → staff receipt.',.65,5.35,4.15)
image(s,A/'admin.png',5.14,2.05,7.5,4.75)
notes(s,'Central NGO and driver face','The central NGO uses its existing volunteer/driver operation and truck fleet. It provisions driver access, maintains equipment and availability, approves facilities and assigns resources. Vehicle capacity and every handling type must match the batch. A driver has one active assignment per resource, records pickup and arrival, and links a unique sticker to the batch. Center staff complete delivery. Map routes are planning visualizations, not live traffic optimization or continuous tracking. In presentation mode, distances and GPS verification are explicitly simulated so the demonstration works on a desktop.')

# 11 · Agentic voice
s = base('Optional, natural-language operation', 'Speak the intent.\nLet the interface handle the steps.',True,35)
shape(s,.68,2.57,7.67,1.02,'ink',True)
text(s,'“I can’t eat cold food. Avoid peanuts.”',.99,2.87,7.02,.46,23,'white')
for i,(name,job) in enumerate([('ElevenLabs','Speech → text'),('Gemini','Correct → interpret → choose an action'),('ElevenLabs','Text → spoken reply')]):
    y=3.98+i*.7
    text(s,name,.75,y,1.9,.34,18,'lime',True)
    text(s,job,2.91,y,5.3,.38,18,'white')
    if i<2: rule(s,.76,y+.5,8.25,y+.5,'teal')
image(s,A/'voice-mobile.png',9.36,1.47,2.41,5.22,False)
foot(s,'Optional alongside manual controls. “Leave a review” opens the sequence; “Submit my review” sends the ready draft once.',True)
notes(s,'ElevenLabs and Gemini assistant','The assistant is an optional agentic layer, not a replacement for the existing interface. Stop recording to send automatically. ElevenLabs transcribes the audio; Gemini lightly corrects transcription errors while preserving intent and negation, then chooses a constrained page action. ElevenLabs speaks the result. A request to leave a review finishes the current visit and opens feedback; a ready review submits when explicitly requested, without repeated confirmation questions. Filters retain visible confirmation, permissions stay with the browser and unrelated requests receive a scoped refusal with available page actions. Audio and text go to their respective providers; keys remain server-side.')

# 12 · Traceability
s = base('Accountability at every handoff', 'An append-style QR evidence trail.',True,36)
steps=[('1','Donor\nphoto'),('2','Event\nQR'),('3','Pickup\nphoto'),('4','Container\nsticker'),('5','Venue\nQR'),('6','Receiving\nphoto')]
rule(s,1.45,3.02,11.85,3.02,'teal',4)
for i,(n,label) in enumerate(steps):
    x=.76+i*2.08
    shape(s,x+.43,2.59,.86,.86,'fern',False,kind=MSO_SHAPE.OVAL)
    text(s,n,x+.43,2.73,.86,.5,28,'white',True,align=PP_ALIGN.CENTER)
    text(s,label,x,3.92,1.74,.9,22,'white',True,align=PP_ALIGN.CENTER)
for i,(title,sub) in enumerate([('QR + timestamps','Stage order and ownership'),('Photos + SHA-256','Evidence linked to the batch'),('Location provenance','Device / simulated / unspecified')]):
    x=.7+i*4.16
    rule(s,x,5.34,x+3.75,5.34,'fern',2)
    text(s,title,x,5.63,3.8,.35,18,'lime',True)
    text(s,sub,x,6.14,3.8,.34,14,'muted')
foot(s,'Traceability supports human review; it is not an immutable ledger or food-safety certificate. Demo GPS is clearly simulated.',True)
notes(s,'QR, images and location evidence','Each valid handoff grows the custody record with a QR event, timestamp, actor and linked photo evidence. The workflow enforces stage order, event and destination identity, assigned-driver permissions, sticker uniqueness and center receipt. SHA-256 hashes identify the stored photo bytes; they do not prove the food is safe or make the whole database immutable. The accurate description is an append-style evidence trail rather than an independently tamper-proof append-only ledger. Live mode has proximity checks; the current desktop demo uses sample coordinates, skips proximity rejection and permanently marks location evidence simulated. AI uses this context for consistency review, not guaranteed fraud detection.')

# 13 · Private reviews
s = base('Learning without public shaming', 'Private feedback. Human guidance.')
shape(s,.66,2.32,4.46,3.64,'white',True,'pale')
text(s,'EXAMPLE FEEDBACK',.96,2.68,3.8,.25,11,'teal',True)
text(s,'“I waited a long time\nand did not receive\na meal.”',.97,3.27,3.81,1.66,27,'ink',font=HEAD)
text(s,'A person’s experience matters.',.97,5.31,3.8,.35,16,'iron')
point(s,'01','Gemini assists the review.','Summarize concerns, flag patterns\nand assess evidence consistency.',5.71,2.43,6.15)
point(s,'02','The central NGO sees it.','Private operational learning,\nnot public ratings or venue shaming.',5.71,3.91,6.15)
point(s,'03','People decide the response.','Follow up with donors and venues;\nuse uncertainty to guide investigation.',5.71,5.39,6.15)
foot(s,'A finish-button interval is approximate. AI cannot verify truthfulness, actual queue time or food safety.')
notes(s,'Private feedback and review verification','Consumer and staff feedback is private to the central NGO. Gemini can identify concerns such as a long wait, packaging problems or temperature worries, summarize them and assess consistency with available evidence. This is review verification in the limited sense of evidence consistency, not proof that a statement is true or false. A long interval between check-in and pressing Finish is not a measured queue duration. Simulated location provides no attendance corroboration. Staff make the judgment and guide donors or venues. Visitor AI results do not automatically impose public ratings or penalties.')

# 14 · Whole workflow and honest boundaries
s = base('From surplus to a shared table', 'One coordinated operational journey.')
journey=[('PLAN','Verified event\nand menu'),('OFFER','Leftovers +\nhandling types'),('DISPATCH','Driver, vehicle\nand destination'),('HAND OFF','QR + photo\ncustody stages'),('RECEIVE','Venue staff\nconfirm intake')]
for i,(title,body) in enumerate(journey):
    x=.65+i*2.5
    shape(s,x,2.62,2.18,1.74,'white',True,'pale')
    text(s,title,x+.17,2.94,1.88,.28,13,'teal',True)
    text(s,body,x+.17,3.47,1.91,.68,17,'ink')
    if i<4: text(s,'→',x+2.22,3.16,.3,.5,24,'fern')
rule(s,.68,4.87,12.62,4.87,'pale')
text(s,'PUBLIC EXPERIENCE',.7,5.13,3.0,.3,12,'teal',True)
text(s,'Preferences → find a place → venue check-in → finish → private review',.7,5.64,11.8,.56,23,'ink',True)
foot(s,'Current demo: fictional meal offers and simulated GPS. Automatically publishing delivered food to Find a Meal is the next integration.')
notes(s,'Operational workflow and demonstration boundaries','The complete operational path has been exercised: register and approve a venue, confirm a donor event menu, create leftovers after the event starts, assign an eligible driver and vehicle, complete QR/photo handoffs and receive the food. The public visit and private review path also works. Be precise about the gap: delivered batches do not yet automatically publish live meal offers to Find a Meal. The public catalog uses fictional offers alongside a separate directory of sourced NYC locations. The current demo deliberately simulates geolocation. These are explicit prototype boundaries rather than claims of production availability or verified GPS.')

# 15 · Sources before the closing statement
s = base('Evidence behind the story', 'Sources & scope',False,35)
for i,(ref,title,scope,url) in enumerate(SOURCES):
    col=0 if i<5 else 1; row=i if i<5 else i-5
    x=.65+col*6.24; y=2.12+row*.87
    link(s,f'[{ref}] {title}',url,x,y,5.9,13)
    text(s,scope,x,y+.29,5.77,.48,11,'iron')
foot(s,'Sources checked 25 October 2026. Product screenshots: current local Samrit app; menu: fictional OCR test fixture.')
notes(s,'Sources and scope','Every source title on this slide is clickable. The statistics retain their original year, geography and denominator. The smartphone statistic is specifically a lower-income ownership measure; the kitchen-network count combines several provider types. The Sanskrit reference supports nectar; essence is the team’s interpretation. The Good Samaritan Act reference is inspiration and context, not a claim of blanket liability protection. The final statement is grounded in the international right to adequate food. Screenshots show the project itself, and the sample menu is fictional.',[v[0] for v in SOURCES])

# 16 · Required closing
s = base('Our reason for building', '',True)
text(s,'Access to food\nis a basic\nhuman right.',.69,1.48,10.8,3.55,55,'white',font=HEAD)
text(s,'A place at the table.\nA community behind it.',.73,5.58,8.3,.96,23,'muted')
logo(s,10.35,2.55,1.72,leaf=True)
notes(s,'Access to food is a basic human right','Close with the principle that unites the product: access to food is a basic human right. Samrit’s role is to make that access easier to navigate and the supporting operations easier to account for. The person seeking a meal comes first. A welcoming experience and a trustworthy organization belong together.', ['S9'])

# Validate slide geometry before saving. Native PowerPoint rendering is checked separately.
problems=[]
for i,slide in enumerate(P.slides,1):
    for sh in slide.shapes:
        if sh.left < -10 or sh.top < -10 or sh.left+sh.width>P.slide_width+10 or sh.top+sh.height>P.slide_height+10:
            problems.append((i, sh.name, 'off slide'))
assert not problems, problems
target = OUT/'Samrit_Final_Presentation.pptx'
P.save(target)
(OUT/'slide_manifest.json').write_text(json.dumps(INDEX,ensure_ascii=False,indent=2),encoding='utf-8')
speaker='\n\n'.join(f"## {v['slide']:02} · {v['title']}\n\n{v['notes']}" for v in INDEX)
(OUT/'SPEAKER_NOTES.md').write_text('# Samrit — presenter notes\n\n' + speaker + '\n\n## Sources\n\n' + '\n\n'.join(f'[{a}] [{b}]({d}) — {c}' for a,b,c,d in SOURCES),encoding='utf-8')
print(f'Created {target.name}: {len(P.slides)} editable slides with speaker notes.')
